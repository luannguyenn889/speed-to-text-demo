import os
import glob
import time
import re
import warnings
import docx
import chromadb
from sentence_transformers import SentenceTransformer

# Ẩn cảnh báo nhắc token của Hugging Face
warnings.filterwarnings("ignore", message=".*unauthenticated requests.*")
warnings.filterwarnings("ignore", message=".*HF Hub.*")
os.environ["HF_HUB_DISABLE_SYMLINKS_WARNING"] = "1"
os.environ["TOKENIZERS_PARALLELISM"] = "false"

# Thư mục chứa tài liệu và cơ sở dữ liệu vector
DOCS_DIR = os.path.join(os.path.dirname(__file__), "document")
DB_DIR = os.path.join(os.path.dirname(__file__), "chroma_db")
COLLECTION_NAME = "qnu_viva_documents"
EMBEDDING_MODEL_NAME = "paraphrase-multilingual-MiniLM-L12-v2"
RERANKER_MODEL_NAME = "cross-encoder/ms-marco-MiniLM-L-6-v2"

_embedding_model = None
_reranker_model = None
_chroma_client = None
_collection = None
_bm25_index = None
_bm25_docs = None

# Từ ngữ đệm / hội thoại thi vấn đáp cần làm sạch trước khi truy vấn véc-tơ
FILLER_PATTERNS = [
    r'\bdạ\b', r'\bthưa thầy\b', r'\bthưa cô\b', r'\bem xin phép\b',
    r'\btheo em nghĩ\b', r'\btheo em nhớ\b', r'\bem nghĩ là\b',
    r'\bừm\b', r'\bà\b', r'\bthì là mà\b', r'\bạ\b'
]

def clean_oral_query(text):
    """Lọc bỏ các từ đệm, ngập ngừng hội thoại để làm sắc nét véc-tơ truy vấn"""
    if not text:
        return ""
    cleaned = text
    for pat in FILLER_PATTERNS:
        cleaned = re.sub(pat, ' ', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'\s+', ' ', cleaned).strip()
    return cleaned if len(cleaned) > 5 else text

def tokenize_text(text):
    """Tách từ đơn giản cho BM25 (hỗ trợ cả tiếng Việt và thuật ngữ CNTT)"""
    if not text:
        return []
    return re.findall(r'\b\w+\b', text.lower())

def get_embedding_model():
    """Tải embedding model theo cơ chế Lazy load từ cache máy tính"""
    global _embedding_model
    if _embedding_model is None:
        print(f"📦 Đang tải Embedding Model ({EMBEDDING_MODEL_NAME})...")
        try:
            _embedding_model = SentenceTransformer(EMBEDDING_MODEL_NAME, local_files_only=True)
        except Exception:
            _embedding_model = SentenceTransformer(EMBEDDING_MODEL_NAME)
        print("✅ Đã nạp Embedding Model thành công.")
    return _embedding_model

def get_reranker_model():
    """Tải Cross-Encoder Reranker model để tái xếp hạng Top-K ứng viên"""
    global _reranker_model
    if _reranker_model is None:
        try:
            from sentence_transformers import CrossEncoder
            try:
                _reranker_model = CrossEncoder(RERANKER_MODEL_NAME, local_files_only=True)
            except Exception:
                _reranker_model = CrossEncoder(RERANKER_MODEL_NAME)
            print(f"✅ Đã nạp Reranker Model ({RERANKER_MODEL_NAME}) thành công.")
        except Exception as e:
            print(f"⚠️ Không thể nạp Reranker ({e}). Hệ thống sẽ dùng điểm Hybrid RRF thay thế.")
            _reranker_model = None
    return _reranker_model

def get_chroma_collection():
    """Khởi tạo ChromaDB client và collection"""
    global _chroma_client, _collection
    if _collection is None:
        os.makedirs(DB_DIR, exist_ok=True)
        _chroma_client = chromadb.PersistentClient(path=DB_DIR)
        _collection = _chroma_client.get_or_create_collection(
            name=COLLECTION_NAME,
            metadata={"hnsw:space": "cosine"}
        )
    return _collection

def get_bm25_index():
    """Khởi tạo hoặc lấy chỉ mục BM25 từ toàn bộ tài liệu trong ChromaDB"""
    global _bm25_index, _bm25_docs
    if _bm25_index is None:
        collection = get_chroma_collection()
        if collection.count() == 0:
            index_documents(force_reindex=False)
            
        total_count = collection.count()
        _bm25_docs = []
        batch_limit = 2000
        for offset in range(0, total_count, batch_limit):
            batch_data = collection.get(
                include=["documents", "metadatas"],
                limit=batch_limit,
                offset=offset
            )
            if batch_data and batch_data.get("documents"):
                for doc, meta in zip(batch_data["documents"], batch_data["metadatas"]):
                    _bm25_docs.append({
                        "text": doc,
                        "source": meta.get("source", "Không rõ"),
                        "index": meta.get("index", 0)
                    })
            
            try:
                from rank_bm25 import BM25Okapi
                tokenized_corpus = [tokenize_text(d["text"]) for d in _bm25_docs]
                _bm25_index = BM25Okapi(tokenized_corpus)
                print(f"✅ Đã khởi tạo chỉ mục BM25 ({len(_bm25_docs)} đoạn tài liệu).")
            except Exception as e:
                print(f"⚠️ Lỗi khởi tạo BM25: {str(e)}")
                _bm25_index = None
    return _bm25_index, _bm25_docs

def extract_text_from_docx(file_path):
    """Trích xuất toàn bộ văn bản và bảng biểu từ file .docx"""
    try:
        doc = docx.Document(file_path)
        content_parts = []
        
        # Đọc từng đoạn văn
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                content_parts.append(text)
                
        # Đọc từng bảng biểu
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                unique_cells = []
                for c in row_cells:
                    if not unique_cells or c != unique_cells[-1]:
                        unique_cells.append(c)
                if unique_cells:
                    content_parts.append(" | ".join(unique_cells))
                    
        return "\n".join(content_parts)
    except Exception as e:
        print(f"⚠️ Lỗi khi đọc file {file_path}: {str(e)}")
        return ""

def extract_text_from_txt(file_path):
    """Trích xuất văn bản từ file text thuần"""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
    except Exception as e:
        print(f"⚠️ Lỗi khi đọc file text {file_path}: {str(e)}")
        return ""

def extract_text_from_pdf(file_path):
    """Trích xuất toàn bộ văn bản từ file .pdf (hỗ trợ cả PyMuPDF fitz và pypdf)"""
    # Cách 1: Ưu tiên PyMuPDF (fitz) vì tốc độ cao và giữ trọn vẹn dấu tiếng Việt
    try:
        import fitz
        doc = fitz.open(file_path)
        pages_text = []
        for page in doc:
            t = page.get_text().strip()
            if t:
                pages_text.append(t)
        doc.close()
        if pages_text:
            return "\n\n".join(pages_text)
    except Exception:
        pass

    # Cách 2: Giải pháp dự phòng bằng pypdf
    try:
        import pypdf
        reader = pypdf.PdfReader(file_path)
        pages_text = []
        for page in reader.pages:
            t = page.extract_text()
            if t and t.strip():
                pages_text.append(t.strip())
        if pages_text:
            return "\n\n".join(pages_text)
    except Exception as e:
        print(f"⚠️ Lỗi khi đọc file PDF {file_path}: {str(e)}")

    return ""

def get_clean_overlap(text, max_chars=100):
    """Lấy phần đuôi để overlap nhưng bắt buộc giữ trọn vẹn từ ngữ (không cắt ngang từ)"""
    if not text or len(text) <= max_chars:
        return text
    tail = text[-max_chars:].strip()
    first_space = tail.find(" ")
    if first_space != -1 and first_space < len(tail) - 1:
        cleaned = tail[first_space + 1:].strip()
        cleaned = cleaned.lstrip(".,:;|-)]} ")
        return cleaned
    return tail

def chunk_text(text, source_name, chunk_size=600, overlap=100):
    """
    Chia văn bản thành các đoạn nhỏ (chunks) theo ranh giới đoạn văn và câu hoàn chỉnh.
    Đảm bảo tuyệt đối không bị cắt cụt từ ngữ ở đầu hoặc cuối mỗi đoạn trích xuất.
    """
    if not text:
        return []
        
    paragraphs = text.split("\n")
    chunks = []
    current_chunk = ""
    chunk_index = 1
    
    for p in paragraphs:
        p = p.strip()
        if not p:
            continue
            
        if len(current_chunk) + len(p) + 1 <= chunk_size:
            current_chunk = f"{current_chunk}\n{p}" if current_chunk else p
        else:
            if current_chunk:
                clean_entry = current_chunk.strip().lstrip(".,:;|-)]} ")
                if len(clean_entry) > 30:
                    chunks.append({
                        "id": f"{source_name}_chunk_{chunk_index}",
                        "text": clean_entry,
                        "source": source_name,
                        "index": chunk_index
                    })
                    chunk_index += 1
                
                tail_overlap = get_clean_overlap(current_chunk, max_chars=overlap)
                current_chunk = f"{tail_overlap}\n{p}" if tail_overlap else p
            else:
                import re
                sentences = re.split(r'(?<=[.?!])\s+', p)
                sub_chunk = ""
                for s in sentences:
                    if len(sub_chunk) + len(s) + 1 <= chunk_size:
                        sub_chunk = f"{sub_chunk} {s}".strip()
                    else:
                        if sub_chunk:
                            clean_s = sub_chunk.strip().lstrip(".,:;|-)]} ")
                            chunks.append({
                                "id": f"{source_name}_chunk_{chunk_index}",
                                "text": clean_s,
                                "source": source_name,
                                "index": chunk_index
                            })
                            chunk_index += 1
                        sub_chunk = s
                current_chunk = sub_chunk
                
    if current_chunk.strip():
        clean_entry = current_chunk.strip().lstrip(".,:;|-)]} ")
        if len(clean_entry) > 20:
            chunks.append({
                "id": f"{source_name}_chunk_{chunk_index}",
                "text": clean_entry,
                "source": source_name,
                "index": chunk_index
            })
        
    return chunks

def index_documents(force_reindex=False):
    """
    Quét toàn bộ tài liệu trong thư mục document/ và nạp vào ChromaDB.
    Đồng thời reset chỉ mục BM25 để cập nhật dữ liệu mới.
    """
    global _bm25_index, _bm25_docs
    collection = get_chroma_collection()
    existing_count = collection.count()
    
    if existing_count > 0 and not force_reindex:
        print(f"ℹ️ Cơ sở dữ liệu RAG đã có sẵn {existing_count} đoạn trích xuất. Bỏ qua bước index.")
        get_bm25_index()  # Đảm bảo BM25 được nạp
        return existing_count
        
    if force_reindex and existing_count > 0:
        print("🔄 Đang xóa index cũ để lập lại chỉ mục mới...")
        global _chroma_client, _collection
        _chroma_client.delete_collection(COLLECTION_NAME)
        _collection = _chroma_client.get_or_create_collection(
            name=COLLECTION_NAME,
            metadata={"hnsw:space": "cosine"}
        )
        collection = _collection
        _bm25_index = None
        _bm25_docs = None
        
    if not os.path.exists(DOCS_DIR):
        os.makedirs(DOCS_DIR, exist_ok=True)
        return 0
        
    docx_files = glob.glob(os.path.join(DOCS_DIR, "*.docx"))
    txt_files = glob.glob(os.path.join(DOCS_DIR, "*.txt"))
    pdf_files = glob.glob(os.path.join(DOCS_DIR, "*.pdf"))
    all_files = docx_files + txt_files + pdf_files
    
    if not all_files:
        print("⚠️ Thư mục document/ hiện chưa có file tài liệu nào.")
        return 0
        
    all_chunks = []
    print(f"📂 Tìm thấy {len(all_files)} tài liệu (.docx, .txt, .pdf). Bắt đầu trích xuất...")
    
    for file_path in all_files:
        base_name = os.path.basename(file_path)
        if base_name.startswith("~$"):
            continue
            
        print(f"  -> Đang đọc: {base_name}")
        if file_path.endswith(".docx"):
            raw_text = extract_text_from_docx(file_path)
        elif file_path.endswith(".pdf"):
            raw_text = extract_text_from_pdf(file_path)
        else:
            raw_text = extract_text_from_txt(file_path)
            
        file_chunks = chunk_text(raw_text, source_name=base_name)
        all_chunks.extend(file_chunks)
        print(f"     Tạo được {len(file_chunks)} đoạn text.")
        
    if not all_chunks:
        print("⚠️ Không có nội dung nào được trích xuất từ tài liệu.")
        return 0
        
    print(f"🚀 Đang tạo Vector Embedding cho tổng cộng {len(all_chunks)} đoạn...")
    model = get_embedding_model()
    
    texts = [c["text"] for c in all_chunks]
    ids = [c["id"] for c in all_chunks]
    metadatas = [{"source": c["source"], "index": c["index"]} for c in all_chunks]
    
    embeddings = model.encode(texts, batch_size=32, show_progress_bar=True, normalize_embeddings=True)
    
    # Nạp vào ChromaDB theo từng lô nhỏ (ChromaDB giới hạn max batch size là 5461)
    CHROMA_BATCH_SIZE = 1000
    for i in range(0, len(all_chunks), CHROMA_BATCH_SIZE):
        batch_end = min(i + CHROMA_BATCH_SIZE, len(all_chunks))
        collection.add(
            ids=ids[i:batch_end],
            documents=texts[i:batch_end],
            metadatas=metadatas[i:batch_end],
            embeddings=embeddings[i:batch_end].tolist()
        )
    
    # Khởi tạo lại BM25
    _bm25_index = None
    _bm25_docs = None
    get_bm25_index()
    
    print(f"✅ Đã hoàn tất lập chỉ mục {len(all_chunks)} đoạn vào ChromaDB & BM25!")
    return len(all_chunks)

def query_relevant_chunks(query_text, top_k=3, score_threshold=0.85, search_mode="hybrid", use_rerank=True):
    """
    Truy vấn nâng cao các đoạn giáo trình liên quan nhất (Advanced Two-Stage Retrieval):
    - Giai đoạn 1: Tìm kiếm kết hợp Hybrid (Dense Vector HNSW + Sparse BM25) qua thuật toán RRF.
    - Giai đoạn 2: Tái xếp hạng (Cross-Encoder Re-ranking) để tối đa hóa độ chính xác ngữ cảnh.
    
    Tham số:
      - query_text: Câu hỏi thi / câu trả lời sinh viên
      - top_k: Số đoạn tài liệu tối ưu trả về (mặc định 3)
      - search_mode: 'hybrid' (kết hợp), 'dense' (chỉ vector), hoặc 'bm25' (chỉ từ khóa)
      - use_rerank: Bật/tắt tầng Cross-Encoder Reranker
    """
    if not query_text or not query_text.strip():
        return []
        
    collection = get_chroma_collection()
    if collection.count() == 0:
        index_documents(force_reindex=False)
        
    if collection.count() == 0:
        return []

    # Tiền xử lý làm sạch từ đệm hội thoại
    cleaned_query = clean_oral_query(query_text)
    candidate_pool_size = max(top_k * 3, 8)  # Lấy nhóm ứng viên rộng hơn trước khi Rerank
    candidate_pool_size = min(candidate_pool_size, collection.count())

    candidates = {}  # key: doc_text, value: dict thông tin

    # 1. TÌM KIẾM BẰNG DENSE VECTOR (CHROMADB)
    if search_mode in ["hybrid", "dense"]:
        model = get_embedding_model()
        query_vector = model.encode([cleaned_query], normalize_embeddings=True).tolist()
        
        results = collection.query(
            query_embeddings=query_vector,
            n_results=candidate_pool_size
        )
        
        if results and "documents" in results and results["documents"]:
            docs = results["documents"][0]
            metas = results["metadatas"][0] if "metadatas" in results else [{}] * len(docs)
            distances = results["distances"][0] if "distances" in results else [0.0] * len(docs)
            
            for rank, (doc, meta, dist) in enumerate(zip(docs, metas, distances), 1):
                sim = round(1.0 - dist, 4)
                if doc not in candidates:
                    candidates[doc] = {
                        "text": doc,
                        "source": meta.get("source", "Không rõ"),
                        "index": meta.get("index", 0),
                        "dense_sim": sim,
                        "dense_rank": rank,
                        "bm25_rank": 999,
                        "rrf_score": 0.0
                    }
                else:
                    candidates[doc]["dense_sim"] = sim
                    candidates[doc]["dense_rank"] = rank

    # 2. TÌM KIẾM BẰNG SPARSE BM25 (TỪ KHÓA VIẾT TẮT / CHUYÊN NGÀNH)
    if search_mode in ["hybrid", "bm25"]:
        bm25, bm25_docs = get_bm25_index()
        if bm25 and bm25_docs:
            tokenized_q = tokenize_text(cleaned_query)
            if tokenized_q:
                bm25_scores = bm25.get_scores(tokenized_q)
                # Lấy top các chỉ số có điểm BM25 cao nhất
                top_bm25_indices = sorted(range(len(bm25_scores)), key=lambda i: bm25_scores[i], reverse=True)[:candidate_pool_size]
                
                for rank, idx in enumerate(top_bm25_indices, 1):
                    score = float(bm25_scores[idx])
                    if score <= 0:
                        continue
                    b_doc = bm25_docs[idx]
                    doc_text = b_doc["text"]
                    
                    if doc_text not in candidates:
                        candidates[doc_text] = {
                            "text": doc_text,
                            "source": b_doc["source"],
                            "index": b_doc["index"],
                            "dense_sim": 0.0,
                            "dense_rank": 999,
                            "bm25_score": round(score, 2),
                            "bm25_rank": rank,
                            "rrf_score": 0.0
                        }
                    else:
                        candidates[doc_text]["bm25_score"] = round(score, 2)
                        candidates[doc_text]["bm25_rank"] = rank

    if not candidates:
        return []

    # 3. KẾT HỢP ĐIỂM SỐ BẰNG THUẬT TOÁN RRF (RECIPROCAL RANK FUSION)
    candidate_list = list(candidates.values())
    k_constant = 60.0
    for c in candidate_list:
        score_dense = 1.0 / (k_constant + c.get("dense_rank", 999)) if c.get("dense_rank", 999) < 999 else 0.0
        score_bm25 = 1.0 / (k_constant + c.get("bm25_rank", 999)) if c.get("bm25_rank", 999) < 999 else 0.0
        c["rrf_score"] = round(score_dense + score_bm25, 5)

    # Sắp xếp theo RRF ban đầu
    candidate_list.sort(key=lambda x: x["rrf_score"], reverse=True)

    # 4. GIAI ĐOẠN 2: TÁI XẾP HẠNG (CROSS-ENCODER RE-RANKING)
    if use_rerank and len(candidate_list) > 1:
        reranker = get_reranker_model()
        if reranker is not None:
            pairs = [[cleaned_query, c["text"]] for c in candidate_list]
            try:
                rerank_scores = reranker.predict(pairs)
                for c, r_score in zip(candidate_list, rerank_scores):
                    c["rerank_score"] = round(float(r_score), 4)
                    # Chuẩn hóa về thang similarity [0.0 - 1.0] dùng hàm sigmoid
                    import math
                    sig_score = round(1.0 / (1.0 + math.exp(-float(r_score))), 3)
                    c["similarity"] = sig_score
                    c["method"] = "Hybrid + Cross-Encoder Rerank"
                
                # Sắp xếp lại theo điểm Cross-Encoder
                candidate_list.sort(key=lambda x: x["rerank_score"], reverse=True)
            except Exception as e:
                print(f"⚠️ Lỗi dự đoán Reranker: {e}. Sử dụng điểm RRF.")
                for c in candidate_list:
                    c["similarity"] = c.get("dense_sim", 0.0)
                    c["method"] = "Hybrid RRF"
        else:
            for c in candidate_list:
                c["similarity"] = c.get("dense_sim", 0.0)
                c["method"] = "Hybrid RRF"
    else:
        for c in candidate_list:
            c["similarity"] = c.get("dense_sim", 0.0)
            c["method"] = "Dense / BM25"

    return candidate_list[:top_k]

def format_rag_context(chunks):
    """
    Định dạng danh sách các đoạn trích xuất thành văn bản sạch sẽ để chèn vào Prompt của LLM.
    Bao gồm tên nguồn tài liệu và độ khớp học thuật.
    """
    if not chunks:
        return ""
        
    formatted = []
    for i, c in enumerate(chunks, 1):
        method_str = f" | Phương thức: {c.get('method', 'RAG')}" if 'method' in c else ""
        sim_str = f" (Độ khớp: {c.get('similarity', 'N/A')}{method_str})"
        formatted.append(
            f"[Đoạn trích {i} - Nguồn: {c['source']}{sim_str}]\n{c['text']}"
        )
    return "\n\n---\n\n".join(formatted)

def get_rag_status():
    """Lấy thông tin trạng thái toàn diện của cơ sở tri thức RAG nâng cao"""
    try:
        collection = get_chroma_collection()
        count = collection.count()
        
        docx_files = [os.path.basename(f) for f in glob.glob(os.path.join(DOCS_DIR, "*.docx")) if not os.path.basename(f).startswith("~$")]
        txt_files = [os.path.basename(f) for f in glob.glob(os.path.join(DOCS_DIR, "*.txt"))]
        pdf_files = [os.path.basename(f) for f in glob.glob(os.path.join(DOCS_DIR, "*.pdf"))]
        files = docx_files + txt_files + pdf_files
        
        bm25, bm25_docs = get_bm25_index()
        has_bm25 = bm25 is not None
        
        return {
            "status": "ready" if count > 0 else "empty",
            "total_chunks": count,
            "total_files": len(files),
            "files": files,
            "embedding_model": EMBEDDING_MODEL_NAME,
            "reranker_model": RERANKER_MODEL_NAME,
            "hybrid_bm25_active": has_bm25,
            "features": [
                "Boundary-Preserved Chunking (600 chars, 100 overlap)",
                "ChromaDB HNSW Dense Cosine Search",
                "BM25 Okapi Sparse Keyword Search",
                "Reciprocal Rank Fusion (RRF)",
                "Cross-Encoder Semantic Re-ranking"
            ]
        }
    except Exception as e:
        return {
            "status": "error",
            "error": str(e),
            "total_chunks": 0,
            "total_files": 0,
            "files": []
        }
