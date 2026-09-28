import os
import sys
import time
import argparse
import json
import difflib
import pandas as pd
from datetime import datetime
from dotenv import load_dotenv

load_dotenv()

# Import các hàm từ pipeline_core
from pipeline_core import (
    load_questions_from_file,
    transcribe_audio,
    transcribe_with_phowhisper,
    transcribe_with_gemini_cloud,
    transcribe_with_groq,
    correct_transcript_with_gemini,
    grade_with_gemini,
    unload_whisper_models,
    get_device,
    get_audio_duration
)

def run_benchmark(
    audio_path,
    question,
    rubric,
    stt_models=None,
    llm_models=None,
    enable_correction=True,
    gemini_api_key=None,
    groq_api_key=None,
    output_excel_path=None
):
    if not os.path.exists(audio_path):
        raise FileNotFoundError(f"Không tìm thấy file âm thanh/video tại: {audio_path}")

    gemini_api_key = gemini_api_key or os.environ.get("GEMINI_API_KEY", "")
    groq_api_key = groq_api_key or os.environ.get("GROQ_API_KEY", "")

    if not stt_models:
        stt_models = [
            "base",
            "small",
            "PhoWhisper-small (Local - Tối ưu Tiếng Việt)",
            "Groq API (Cloud - whisper-large-v3)",
            "Gemini API (Cloud - Siêu Nhanh)"
        ]

    if not llm_models:
        llm_models = [
            "gemini-2.5-flash",
            "gemini-3.5-flash",
            "gemini-flash-latest"
        ]

    # Tính độ dài file âm thanh
    audio_duration = get_audio_duration(audio_path)

    print("=" * 70)
    print("🚀 BẮT ĐẦU CHẠY BENCHMARK PIPELINE ĐA MÔ HÌNH VÀ XUẤT EXCEL")
    print(f"📁 File âm thanh: {audio_path} ({audio_duration}s)")
    print(f"📝 Câu hỏi: {question[:80]}...")
    print(f"🎯 Mô hình STT thử nghiệm: {stt_models}")
    print(f"🤖 Mô hình LLM thử nghiệm: {llm_models}")
    print("=" * 70)

    records = []

    for stt_model in stt_models:
        print(f"\n🎙️ [STT] Đang chạy mô hình: {stt_model}...")
        stt_start = time.time()
        raw_transcript = ""
        stt_error = ""

        try:
            if stt_model == "Gemini API (Cloud - Siêu Nhanh)":
                unload_whisper_models()
                raw_transcript = transcribe_with_gemini_cloud(audio_path, gemini_api_key, "gemini-2.5-flash")
            elif stt_model == "Groq API (Cloud - whisper-large-v3)":
                unload_whisper_models()
                raw_transcript = transcribe_with_groq(audio_path, groq_api_key)
            elif "PhoWhisper" in stt_model:
                raw_transcript = transcribe_with_phowhisper(audio_path, "vinai/PhoWhisper-small")
            else:
                raw_transcript = transcribe_audio(audio_path, stt_model)
            stt_duration = round(time.time() - stt_start, 2)
            print(f"   ✅ STT thành công ({stt_duration}s): {raw_transcript[:70]}...")
        except Exception as e:
            stt_duration = round(time.time() - stt_start, 2)
            stt_error = str(e)
            print(f"   ❌ STT thất bại ({stt_duration}s): {stt_error}")

        # Tính toán chỉ số STT
        rtf = round(stt_duration / audio_duration, 3) if audio_duration > 0 else 0.0
        raw_word_count = len(raw_transcript.split()) if raw_transcript else 0
        wpm = round((raw_word_count / (audio_duration / 60.0)), 1) if (audio_duration > 0 and raw_word_count > 0) else 0.0

        if not raw_transcript:
            records.append({
                "STT Model": stt_model,
                "Thời Gian STT (s)": stt_duration,
                "Thời Lượng Audio (s)": audio_duration,
                "Hệ Số RTF": rtf,
                "Số Từ Ký Âm Thô": 0,
                "Tốc Độ Nói (WPM)": 0,
                "Ký Âm Thô (Raw Transcript)": f"LỖI: {stt_error}",
                "Thời Gian Sửa AI (s)": 0,
                "Ký Âm Sau Sửa AI": "",
                "Số Từ Sau Sửa AI": 0,
                "Độ Tương Đồng Sửa (%)": 0,
                "LLM Chấm Điểm": "N/A",
                "Thời Gian Chấm (s)": 0,
                "Tổng Thời Gian Pipeline (s)": stt_duration,
                "Điểm Số (/10)": "N/A",
                "Nhận Xét Chung": f"Không có văn bản do lỗi STT: {stt_error}",
                "Điểm Mạnh": "",
                "Điểm Yếu / Thiếu Sót": "",
                "Gợi Ý Cải Thiện": ""
            })
            continue

        # Sửa lỗi chính tả bằng AI
        corrected_transcript = raw_transcript
        corr_duration = 0.0
        if enable_correction:
            print("   ✨ [AI Correction] Đang chuẩn hóa ngữ cảnh CNTT...")
            corr_start = time.time()
            try:
                corrected_transcript = correct_transcript_with_gemini(
                    raw_transcript, question, rubric, gemini_api_key, "gemini-2.5-flash"
                )
                corr_duration = round(time.time() - corr_start, 2)
            except Exception as e:
                corr_duration = round(time.time() - corr_start, 2)
                print(f"   ⚠️ Lỗi sửa chính tả AI: {str(e)}")
                corrected_transcript = raw_transcript

        corr_word_count = len(corrected_transcript.split()) if corrected_transcript else 0
        similarity_pct = round(difflib.SequenceMatcher(None, raw_transcript, corrected_transcript).ratio() * 100, 1) if raw_transcript else 100.0

        # Chấm điểm với từng LLM model
        for llm_model in llm_models:
            print(f"   🤖 [LLM Grading] Đang chấm điểm với: {llm_model}...")
            llm_start = time.time()
            try:
                evaluation = grade_with_gemini(question, rubric, corrected_transcript, gemini_api_key, llm_model)
                llm_duration = round(time.time() - llm_start, 2)

                score = evaluation.get("score", "N/A")
                summary = evaluation.get("summary", "")
                strengths = "\n• ".join(evaluation.get("strengths", []))
                if strengths: strengths = "• " + strengths
                weaknesses = "\n• ".join(evaluation.get("weaknesses", []))
                if weaknesses: weaknesses = "• " + weaknesses
                improvement = "\n• ".join(evaluation.get("improvement", []))
                if improvement: improvement = "• " + improvement

                total_time = round(stt_duration + corr_duration + llm_duration, 2)

                print(f"      ✅ Chấm điểm xong ({llm_duration}s) -> Điểm: {score}/10")

                records.append({
                    "STT Model": stt_model,
                    "Thời Gian STT (s)": stt_duration,
                    "Thời Lượng Audio (s)": audio_duration,
                    "Hệ Số RTF": rtf,
                    "Số Từ Ký Âm Thô": raw_word_count,
                    "Tốc Độ Nói (WPM)": wpm,
                    "Ký Âm Thô (Raw Transcript)": raw_transcript,
                    "Thời Gian Sửa AI (s)": corr_duration,
                    "Ký Âm Sau Sửa AI": corrected_transcript,
                    "Số Từ Sau Sửa AI": corr_word_count,
                    "Độ Tương Đồng Sửa (%)": similarity_pct,
                    "LLM Chấm Điểm": llm_model,
                    "Thời Gian Chấm (s)": llm_duration,
                    "Tổng Thời Gian Pipeline (s)": total_time,
                    "Điểm Số (/10)": score,
                    "Nhận Xét Chung": summary,
                    "Điểm Mạnh": strengths,
                    "Điểm Yếu / Thiếu Sót": weaknesses,
                    "Gợi Ý Cải Thiện": improvement
                })
            except Exception as e:
                llm_duration = round(time.time() - llm_start, 2)
                total_time = round(stt_duration + corr_duration + llm_duration, 2)
                err_msg = str(e)
                print(f"      ❌ Chấm điểm lỗi ({llm_duration}s): {err_msg}")
                records.append({
                    "STT Model": stt_model,
                    "Thời Gian STT (s)": stt_duration,
                    "Thời Lượng Audio (s)": audio_duration,
                    "Hệ Số RTF": rtf,
                    "Số Từ Ký Âm Thô": raw_word_count,
                    "Tốc Độ Nói (WPM)": wpm,
                    "Ký Âm Thô (Raw Transcript)": raw_transcript,
                    "Thời Gian Sửa AI (s)": corr_duration,
                    "Ký Âm Sau Sửa AI": corrected_transcript,
                    "Số Từ Sau Sửa AI": corr_word_count,
                    "Độ Tương Đồng Sửa (%)": similarity_pct,
                    "LLM Chấm Điểm": llm_model,
                    "Thời Gian Chấm (s)": llm_duration,
                    "Tổng Thời Gian Pipeline (s)": total_time,
                    "Điểm Số (/10)": "LỖI",
                    "Nhận Xét Chung": f"Lỗi chấm điểm: {err_msg}",
                    "Điểm Mạnh": "",
                    "Điểm Yếu / Thiếu Sót": "",
                    "Gợi Ý Cải Thiện": ""
                })

    # Tạo DataFrame kết quả
    df_results = pd.DataFrame(records)

    # Tạo bảng tổng kết thống kê (Summary Sheet)
    stt_summary_records = []
    for model_name, grp in df_results.groupby("STT Model"):
        valid_grp = grp[grp["Ký Âm Thô (Raw Transcript)"].str.startswith("LỖI:") == False]
        if not valid_grp.empty:
            avg_stt_time = round(valid_grp["Thời Gian STT (s)"].mean(), 2)
            avg_rtf = round(valid_grp["Hệ Số RTF"].mean(), 3)
            avg_wpm = round(valid_grp["Tốc Độ Nói (WPM)"].mean(), 1)
            raw_words = int(valid_grp["Số Từ Ký Âm Thô"].iloc[0])
            avg_sim = round(valid_grp["Độ Tương Đồng Sửa (%)"].mean(), 1)
        else:
            avg_stt_time, avg_rtf, avg_wpm, raw_words, avg_sim = 0, 0, 0, 0, 0
        stt_summary_records.append({
            "STT Model": model_name,
            "Thời Gian STT TB (s)": avg_stt_time,
            "Hệ Số RTF (STT/Audio)": avg_rtf,
            "Tốc Độ Nói (WPM)": avg_wpm,
            "Số Từ Ký Âm Được": raw_words,
            "Độ Tương Đồng Sau Sửa (%)": avg_sim
        })
    df_stt_summary = pd.DataFrame(stt_summary_records)

    llm_summary_records = []
    for model_name, grp in df_results.groupby("LLM Chấm Điểm"):
        if model_name == "N/A":
            continue
        valid_scores = pd.to_numeric(grp["Điểm Số (/10)"], errors='coerce').dropna()
        avg_score = round(valid_scores.mean(), 2) if not valid_scores.empty else "N/A"
        avg_llm_time = round(grp["Thời Gian Chấm (s)"].mean(), 2)
        avg_total_time = round(grp["Tổng Thời Gian Pipeline (s)"].mean(), 2)
        llm_summary_records.append({
            "LLM Chấm Điểm": model_name,
            "Điểm Số Trung Bình (/10)": avg_score,
            "Thời Gian Chấm TB (s)": avg_llm_time,
            "Tổng Thời Gian Pipeline TB (s)": avg_total_time
        })
    df_llm_summary = pd.DataFrame(llm_summary_records)

    # Sheet Metadata
    metadata = [
        {"Thuộc Tính": "Thời Gian Thực Hiện", "Giá Trị": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
        {"Thuộc Tính": "File Âm Thanh Test", "Giá Trị": os.path.abspath(audio_path)},
        {"Thuộc Tính": "Thời Lượng File Âm Thanh", "Giá Trị": f"{audio_duration} giây ({round(audio_duration/60, 2)} phút)"},
        {"Thuộc Tính": "Thiết Bị Phần Cứng", "Giá Trị": get_device().upper()},
        {"Thuộc Tính": "Câu Hỏi Đề Thi", "Giá Trị": question},
        {"Thuộc Tính": "Tiêu Chí Rubric", "Giá Trị": rubric},
        {"Thuộc Tính": "Sửa Lỗi Chính Tả AI", "Giá Trị": "Bật (gemini-2.5-flash)" if enable_correction else "Tắt"},
    ]
    df_meta = pd.DataFrame(metadata)

    if not output_excel_path:
        timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        os.makedirs("exports", exist_ok=True)
        output_excel_path = os.path.join("exports", f"benchmark_ket_qua_{timestamp_str}.xlsx")

    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    with pd.ExcelWriter(output_excel_path, engine="openpyxl") as writer:
        df_results.to_excel(writer, sheet_name="KetQuaChiTiet", index=False)
        df_stt_summary.to_excel(writer, sheet_name="TongKetSoSanh", index=False, startrow=1)
        # Ghi bảng LLM tiếp theo ở Sheet TongKetSoSanh
        start_row_llm = len(df_stt_summary) + 4
        df_llm_summary.to_excel(writer, sheet_name="TongKetSoSanh", index=False, startrow=start_row_llm)
        df_meta.to_excel(writer, sheet_name="ThongTinDeThi", index=False)

        # Style Sheets
        header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
        section_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        section_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # 1. Format Sheet KetQuaChiTiet
        ws_results = writer.sheets["KetQuaChiTiet"]
        for col_num in range(1, len(df_results.columns) + 1):
            cell = ws_results.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in ws_results.iter_rows(min_row=2, max_row=len(df_results)+1, min_col=1, max_col=len(df_results.columns)):
            for cell in row:
                cell.border = border_thin
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        col_widths_results = {
            "A": 22, "B": 16, "C": 18, "D": 14, "E": 16, "F": 16, "G": 40,
            "H": 18, "I": 40, "J": 16, "K": 18, "L": 22, "M": 16, "N": 22,
            "O": 14, "P": 45, "Q": 35, "R": 35, "S": 35
        }
        for col_letter, width in col_widths_results.items():
            if col_letter in ws_results.column_dimensions:
                ws_results.column_dimensions[col_letter].width = width

        # 2. Format Sheet TongKetSoSanh
        ws_summary = writer.sheets["TongKetSoSanh"]
        ws_summary.cell(row=1, column=1, value="📊 BẢNG 1: TỔNG KẾT HIỆU NĂNG KÝ ÂM (SPEECH-TO-TEXT)").font = section_font
        ws_summary.cell(row=1, column=1).fill = section_fill

        ws_summary.cell(row=start_row_llm, column=1, value="🤖 BẢNG 2: TỔNG KẾT HIỆU NĂNG CHẤM ĐIỂM (LLM GRADING)").font = section_font
        ws_summary.cell(row=start_row_llm, column=1).fill = section_fill

        for col_letter in ["A", "B", "C", "D", "E", "F"]:
            ws_summary.column_dimensions[col_letter].width = 28

        # 3. Format Sheet ThongTinDeThi
        ws_meta = writer.sheets["ThongTinDeThi"]
        for col_num in range(1, 3):
            cell = ws_meta.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_meta.column_dimensions["A"].width = 26
        ws_meta.column_dimensions["B"].width = 75

        for row in ws_meta.iter_rows(min_row=2, max_row=len(df_meta)+1, min_col=1, max_col=2):
            for cell in row:
                cell.border = border_thin
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    print("\n" + "=" * 70)
    print(f"🎉 ĐÃ XUẤT THÀNH CÔNG BÁO CÁO EXCEL TẠI: {output_excel_path}")
    print("=" * 70)
    return output_excel_path, df_results


def main():
    parser = argparse.ArgumentParser(description="Chạy thử nghiệm đa mô hình STT & LLM và xuất kết quả ra file Excel.")
    parser.add_argument("--audio", type=str, default="audio/caubademo.mp4", help="Đường dẫn file âm thanh/video bài làm.")
    parser.add_argument("--question_idx", type=int, default=1, help="Số thứ tự câu hỏi mẫu trong BỘ ĐỀ THI.txt (mặc định: 1)")
    parser.add_argument("--question", type=str, default=None, help="Nội dung câu hỏi tự nhập (nếu không dùng bộ đề mẫu)")
    parser.add_argument("--rubric", type=str, default=None, help="Nội dung rubric tự nhập")
    parser.add_argument("--output", type=str, default=None, help="Đường dẫn file Excel đầu ra")
    parser.add_argument("--no_correction", action="store_true", help="Tắt bước sửa lỗi chính tả AI")
    args = parser.parse_args()

    presets = load_questions_from_file("BỘ ĐỀ THI.txt")
    preset_keys = list(presets.keys())

    question = args.question
    rubric = args.rubric

    if not question or not rubric:
        if preset_keys and 0 <= args.question_idx - 1 < len(preset_keys):
            selected_key = preset_keys[args.question_idx - 1]
            question = presets[selected_key]["question"]
            rubric = presets[selected_key]["rubric"]
            print(f"📌 Đã nạp câu hỏi mẫu #{args.question_idx}: {selected_key}")
        else:
            question = "Trình bày khái niệm và vai trò của lập trình hướng đối tượng (OOP). Nêu 4 tính chất cơ bản của OOP."
            rubric = "- Trình bày rõ định nghĩa OOP (1.5 điểm)\n- Nêu đúng tên và định nghĩa ngắn gọn 4 tính chất: Đóng gói, Kế thừa, Đa hình, Trừu tượng (Mỗi tính chất 1.5 điểm, tổng 6.0 điểm)\n- Lấy ví dụ minh họa thực tế dễ hiểu (1.5 điểm)\n- Kỹ năng diễn đạt, trả lời lưu loát (1.0 điểm)"

    run_benchmark(
        audio_path=args.audio,
        question=question,
        rubric=rubric,
        enable_correction=not args.no_correction,
        output_excel_path=args.output
    )

if __name__ == "__main__":
    main()
